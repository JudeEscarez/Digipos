"""
Load the Digipos demo workbook, score every lead, and write the results.

    pip install openpyxl
    python scripts/import_xlsx.py path/to/Digipos_demo_data.xlsx            # writes data/seed.json
    python scripts/import_xlsx.py path/to/Digipos_demo_data.xlsx --sql out.sql   # also SQL for Supabase
    python scripts/import_xlsx.py path/to/Digipos_demo_data.xlsx --push      # upsert straight to Supabase
                                                                             # (needs SUPABASE_URL + SUPABASE_SERVICE_ROLE_KEY)

Re-run it whenever the sheet changes. Visits already logged are kept; scores are recalculated.
"""

import argparse
import importlib.util
import json
import os
import sys
from datetime import date, datetime

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
spec = importlib.util.spec_from_file_location("api_core", os.path.join(ROOT, "api", "index.py"))
core = importlib.util.module_from_spec(spec)
spec.loader.exec_module(core)

with open(os.path.join(ROOT, "data", "example_new_openings.json")) as f:
    EXAMPLE_NEW_OPENINGS = json.load(f)


def iso(v):
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    return str(v)[:10] if v else None


def yes(v):
    return str(v or "").strip().upper().startswith("Y")


def num(v, cast=float):
    try:
        return cast(v) if v not in (None, "") else None
    except (TypeError, ValueError):
        return None


def read_leads(wb):
    ws = wb["Leads"]
    header = [str(c.value or "").strip() for c in ws[1]]
    col = {h.split(" (")[0]: i for i, h in enumerate(header)}
    out = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        lid, name = r[col["lead_id"]], r[col["business_name"]]
        if not lid or not name or not str(lid).startswith("L"):
            continue
        lane_label = str(r[col["lane"]] or "Established")
        area, emirate = r[col["area"]], r[col["emirate"]]
        out.append({
            "id": str(lid), "business_name": str(name).strip(),
            "lane": "new_opening" if lane_label.lower().startswith("new") else "established",
            "lane_label": lane_label, "category": r[col["category"]], "emirate": emirate,
            "area": area, "cluster": core.cluster_for(area, emirate),
            "maps_link": r[col["google_maps_link"]], "website": r[col["website_or_social"]],
            "rating": num(r[col["rating"]]), "review_count": num(r[col["review_count"]], int),
            "price_level": num(r[col["price_level"]], int), "has_phone": yes(r[col["has_phone"]]),
            "oldest_review_date": iso(r[col["oldest_review_date"]]),
            "licence_issue_date": None,
            "card_signal": r[col["card_signal"]], "card_signal_note": r[col["card_signal_note"]],
            "restricted": yes(r[col["restricted_flag"]]), "chain_check": r[col["chain_check"]],
            "date_collected": iso(r[col["date_collected"]]),
            "data_source": "google_places", "is_example": False,
        })
    return out


def read_merchants(wb):
    ws = wb["Merchants (simulated)"]
    out = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r[0] or not str(r[0]).startswith("M"):
            continue
        vols = [num(v, int) for v in r[8:14]]
        out.append({
            "id": r[0], "name": r[1], "category": r[2], "emirate": r[3], "area": r[4],
            "activation_date": iso(r[5]), "terminals": num(r[6], int),
            "courtesy_call_done": yes(r[7]),
            "vol_apr": vols[0], "vol_may": vols[1], "vol_jun": vols[2],
            "vol_jul": vols[3], "vol_aug": vols[4], "vol_sep": vols[5], "is_simulated": True,
        })
    return out


def examples():
    out = []
    for e in EXAMPLE_NEW_OPENINGS:
        row = {k: None for k in ("maps_link", "website", "rating", "review_count", "price_level",
                                 "oldest_review_date", "card_signal", "card_signal_note",
                                 "chain_check", "date_collected")}
        row.update(e)
        row.update(lane="new_opening", lane_label="New opening — EXAMPLE FORMAT, not a real business",
                   emirate="Dubai", cluster=core.cluster_for(e["area"], "Dubai"),
                   restricted=False, data_source="example", is_example=True)
        out.append(row)
    return out


def sql_literal(v):
    if v is None:
        return "null"
    if isinstance(v, bool):
        return "true" if v else "false"
    if isinstance(v, (int, float)):
        return str(v)
    if isinstance(v, (dict, list)):
        v = json.dumps(v)
    return "'" + str(v).replace("'", "''") + "'"


def to_sql(table, rows, update_cols):
    cols = list(rows[0].keys())
    values = ",\n".join("(" + ", ".join(sql_literal(r[c]) for c in cols) + ")" for r in rows)
    sets = ", ".join(f"{c} = excluded.{c}" for c in update_cols)
    return (f"insert into public.{table} ({', '.join(cols)}) values\n{values}\n"
            f"on conflict (id) do update set {sets};\n")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx")
    ap.add_argument("--sql", help="write Supabase upsert SQL to this file")
    ap.add_argument("--push", action="store_true", help="upsert into Supabase via the API store")
    ap.add_argument("--no-examples", action="store_true", help="skip the example new-opening rows")
    a = ap.parse_args()

    wb = openpyxl.load_workbook(a.xlsx, data_only=False)
    leads = read_leads(wb) + ([] if a.no_examples else examples())
    merchants = read_merchants(wb)

    core.rescore(leads, [])  # fresh scores, no visits
    for l in leads:
        l["draft_message"] = core.draft_message(l)
        l.setdefault("status", "open")

    with open(os.path.join(ROOT, "data", "seed.json"), "w") as f:
        json.dump({"leads": leads, "merchants": merchants}, f, indent=1, ensure_ascii=False)
    print(f"data/seed.json: {len(leads)} leads, {len(merchants)} merchants")

    if a.sql:
        lead_cols = [c for c in leads[0] if c not in ("id", "status")]
        with open(a.sql, "w") as f:
            f.write(to_sql("leads", leads, lead_cols))
            f.write(to_sql("merchants", merchants, [c for c in merchants[0] if c != "id"]))
        print("SQL written to", a.sql)

    if a.push:
        store = core.get_store()
        if not isinstance(store, core.SupabaseStore):
            sys.exit("Set SUPABASE_URL and SUPABASE_SERVICE_ROLE_KEY to push.")
        store.insert_or_replace("leads", [{k: v for k, v in l.items() if k != "status"} for l in leads])
        store.insert_or_replace("merchants", merchants)
        # Re-apply any logged visits so statuses and coverage survive a re-import.
        store.upsert_leads(core.rescore(store.leads(), store.visits()))
        print("Pushed to Supabase.")


if __name__ == "__main__":
    main()
